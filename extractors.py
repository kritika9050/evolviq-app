from pathlib import Path
from bs4 import BeautifulSoup
import re
from pypdf import PdfReader
from docx import Document
import openpyxl


def clean_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    boilerplate = ["cookie policy", "privacy policy", "subscribe", "all rights reserved"]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    kept = [s for s in sentences if len(s) > 35 and not any(b in s.lower() for b in boilerplate)]
    return " ".join(kept)


def extract_from_file(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(p))
        return clean_text("\n".join(page.extract_text() or "" for page in reader.pages))
    if suffix == ".docx":
        doc = Document(str(p))
        return clean_text("\n".join(par.text for par in doc.paragraphs))
    if suffix == ".xlsx":
        wb = openpyxl.load_workbook(str(p), data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append(" | ".join(str(c) for c in row if c is not None))
        return clean_text("\n".join(rows))
    if suffix in [".txt", ".md"]:
        return clean_text(p.read_text(errors="ignore"))
    return ""


def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return clean_text(soup.get_text(" "))
